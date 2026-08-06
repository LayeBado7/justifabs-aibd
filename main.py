import csv
import hashlib
import io
import os
import secrets
import smtplib
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Optional
from urllib.parse import quote

from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook
from passlib.context import CryptContext
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from sqlalchemy import Boolean, Date, DateTime, ForeignKey, Integer, String, Text, create_engine, func, or_, select
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker

BASE_DIR = Path(__file__).resolve().parent
APP_NAME = os.getenv("APP_NAME", "JustifAbs AIBD SA")
ENVIRONMENT = os.getenv("ENVIRONMENT", "production").lower()
SECRET_KEY = os.getenv("SECRET_KEY", "")
if ENVIRONMENT == "production" and len(SECRET_KEY) < 32:
    raise RuntimeError("SECRET_KEY doit contenir au moins 32 caractères en production")
if not SECRET_KEY:
    SECRET_KEY = secrets.token_urlsafe(48)
DATABASE_URL = os.getenv("DATABASE_URL", f"sqlite:///{BASE_DIR / 'data' / 'aibd_absences.db'}")
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", str(BASE_DIR / "data" / "uploads")))
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "10"))
MAX_FILES_PER_DOSSIER = int(os.getenv("MAX_FILES_PER_DOSSIER", "10"))
ALLOWED_EXTENSIONS = {x.strip().lower() for x in os.getenv("ALLOWED_EXTENSIONS", "pdf,jpg,jpeg,png,doc,docx").split(",")}
COOKIE_SECURE = os.getenv("COOKIE_SECURE", "true").lower() == "true"
SESSION_MAX_AGE = int(os.getenv("SESSION_MAX_AGE", "28800"))
BOOTSTRAP_ADMIN_EMAIL = os.getenv("BOOTSTRAP_ADMIN_EMAIL", "").strip().lower()
BOOTSTRAP_ADMIN_PASSWORD = os.getenv("BOOTSTRAP_ADMIN_PASSWORD", "")
BOOTSTRAP_ADMIN_NAME = os.getenv("BOOTSTRAP_ADMIN_NAME", "Administrateur JustifAbs")
BOOTSTRAP_ADMIN_MATRICULE = os.getenv("BOOTSTRAP_ADMIN_MATRICULE", "ADM001")
SMTP_HOST = os.getenv("SMTP_HOST", "")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER", "")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "")
SMTP_FROM = os.getenv("SMTP_FROM", SMTP_USER)
SMTP_TLS = os.getenv("SMTP_TLS", "true").lower() == "true"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
if DATABASE_URL.startswith("sqlite"):
    db_path = DATABASE_URL.replace("sqlite:///", "", 1)
    if db_path and db_path != ":memory:":
        Path(db_path).parent.mkdir(parents=True, exist_ok=True)

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(DATABASE_URL, connect_args=connect_args, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
serializer = URLSafeTimedSerializer(SECRET_KEY, salt="aibd-session")


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(primary_key=True)
    matricule: Mapped[str] = mapped_column(String(30), unique=True, index=True)
    full_name: Mapped[str] = mapped_column(String(150))
    email: Mapped[str] = mapped_column(String(180), unique=True, index=True)
    department: Mapped[str] = mapped_column(String(150), default="")
    site: Mapped[str] = mapped_column(String(100), default="DIASS")
    role: Mapped[str] = mapped_column(String(20), default="agent")
    password_hash: Mapped[str] = mapped_column(String(255))
    is_active: Mapped[bool] = mapped_column(Boolean, default=True)
    must_change_password: Mapped[bool] = mapped_column(Boolean, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    absences: Mapped[list["Absence"]] = relationship(back_populates="agent", foreign_keys="Absence.agent_id")


class Absence(Base):
    __tablename__ = "absences"
    id: Mapped[int] = mapped_column(primary_key=True)
    agent_id: Mapped[int] = mapped_column(ForeignKey("users.id"), index=True)
    absence_type: Mapped[str] = mapped_column(String(80))
    start_date: Mapped[date] = mapped_column(Date)
    end_date: Mapped[date] = mapped_column(Date)
    reason: Mapped[str] = mapped_column(Text, default="")
    status: Mapped[str] = mapped_column(String(30), default="soumis", index=True)
    submitted_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    validated_at: Mapped[Optional[datetime]] = mapped_column(DateTime, nullable=True)
    validated_by: Mapped[Optional[int]] = mapped_column(ForeignKey("users.id"), nullable=True)
    decision_note: Mapped[str] = mapped_column(Text, default="")
    agent: Mapped[User] = relationship(back_populates="absences", foreign_keys=[agent_id])
    documents: Mapped[list["Document"]] = relationship(back_populates="absence", cascade="all, delete-orphan")
    messages: Mapped[list["Message"]] = relationship(back_populates="absence", cascade="all, delete-orphan", order_by="Message.created_at")
    events: Mapped[list["AuditEvent"]] = relationship(back_populates="absence", cascade="all, delete-orphan", order_by="AuditEvent.created_at")


class Document(Base):
    __tablename__ = "documents"
    id: Mapped[int] = mapped_column(primary_key=True)
    absence_id: Mapped[int] = mapped_column(ForeignKey("absences.id"), index=True)
    original_name: Mapped[str] = mapped_column(String(255))
    stored_name: Mapped[str] = mapped_column(String(255), unique=True)
    content_type: Mapped[str] = mapped_column(String(120), default="application/octet-stream")
    size_bytes: Mapped[int] = mapped_column(Integer)
    sha256: Mapped[str] = mapped_column(String(64))
    uploaded_by: Mapped[int] = mapped_column(ForeignKey("users.id"))
    uploaded_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    absence: Mapped[Absence] = relationship(back_populates="documents")


class Message(Base):
    __tablename__ = "messages"
    id: Mapped[int] = mapped_column(primary_key=True)
    absence_id: Mapped[int] = mapped_column(ForeignKey("absences.id"), index=True)
    author_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    body: Mapped[str] = mapped_column(Text)
    kind: Mapped[str] = mapped_column(String(30), default="commentaire")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    absence: Mapped[Absence] = relationship(back_populates="messages")
    author: Mapped[User] = relationship()


class AuditEvent(Base):
    __tablename__ = "audit_events"
    id: Mapped[int] = mapped_column(primary_key=True)
    absence_id: Mapped[int] = mapped_column(ForeignKey("absences.id"), index=True)
    actor_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    action: Mapped[str] = mapped_column(String(80))
    details: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    absence: Mapped[Absence] = relationship(back_populates="events")
    actor: Mapped[User] = relationship()


Base.metadata.create_all(engine)
app = FastAPI(title=APP_NAME, docs_url=None if ENVIRONMENT == "production" else "/docs")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


def db_session():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def token_data(request: Request) -> dict:
    token = request.cookies.get("session")
    if not token:
        raise HTTPException(401)
    try:
        return serializer.loads(token, max_age=SESSION_MAX_AGE)
    except (BadSignature, SignatureExpired, TypeError, ValueError):
        raise HTTPException(401)


def current_user(request: Request, db: Session = Depends(db_session)) -> User:
    data = token_data(request)
    user = db.get(User, int(data.get("uid", 0)))
    if not user or not user.is_active:
        raise HTTPException(401)
    return user


def optional_user(request: Request, db: Session) -> Optional[User]:
    try:
        data = token_data(request)
        return db.get(User, int(data.get("uid", 0)))
    except Exception:
        return None


def csrf_token(request: Request) -> str:
    try:
        return str(token_data(request).get("csrf", ""))
    except Exception:
        return ""


def verify_csrf(request: Request, submitted: str):
    expected = csrf_token(request)
    if not expected or not secrets.compare_digest(expected, submitted or ""):
        raise HTTPException(403, "Jeton de sécurité invalide. Rechargez la page.")


def context(request: Request, **kwargs):
    return {"request": request, "app_name": APP_NAME, "csrf_token": csrf_token(request), **kwargs}


def flash_redirect(url: str, message: str, level: str = "success"):
    separator = "&" if "?" in url else "?"
    return RedirectResponse(f"{url}{separator}msg={quote(message)}&level={level}", status_code=303)


def can_access(user: User, absence: Absence) -> bool:
    if user.role in {"rh", "admin"}:
        return True
    if user.role == "manager":
        return absence.agent.department == user.department and absence.agent.site == user.site
    return absence.agent_id == user.id


def add_event(db: Session, absence: Absence, actor: User, action: str, details: str = ""):
    db.add(AuditEvent(absence_id=absence.id, actor_id=actor.id, action=action, details=details[:2000]))


def validate_password(password: str):
    if len(password) < 10 or not any(c.isupper() for c in password) or not any(c.islower() for c in password) or not any(c.isdigit() for c in password):
        raise HTTPException(400, "Le mot de passe doit avoir 10 caractères, une majuscule, une minuscule et un chiffre.")


def save_upload(db: Session, absence: Absence, user: User, upload: UploadFile):
    if len(absence.documents) >= MAX_FILES_PER_DOSSIER:
        raise HTTPException(400, f"Maximum {MAX_FILES_PER_DOSSIER} fichiers par dossier")
    original = Path(upload.filename or "document").name[:255]
    ext = original.rsplit(".", 1)[-1].lower() if "." in original else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"Type de fichier non autorisé : .{ext}")
    data = upload.file.read(MAX_UPLOAD_MB * 1024 * 1024 + 1)
    if not data:
        raise HTTPException(400, "Le fichier est vide")
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(400, f"Fichier supérieur à {MAX_UPLOAD_MB} Mo")
    digest = hashlib.sha256(data).hexdigest()
    stored = f"{absence.id}_{secrets.token_hex(16)}.{ext}"
    (UPLOAD_DIR / stored).write_bytes(data)
    db.add(Document(absence_id=absence.id, original_name=original, stored_name=stored,
                    content_type=upload.content_type or "application/octet-stream",
                    size_bytes=len(data), sha256=digest, uploaded_by=user.id))


def send_email(to: str, subject: str, body: str):
    if not SMTP_HOST or not SMTP_FROM or not to:
        return False
    try:
        msg = EmailMessage()
        msg["From"] = SMTP_FROM
        msg["To"] = to
        msg["Subject"] = subject
        msg.set_content(body)
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=10) as server:
            if SMTP_TLS:
                server.starttls()
            if SMTP_USER:
                server.login(SMTP_USER, SMTP_PASSWORD)
            server.send_message(msg)
        return True
    except Exception:
        return False


def bootstrap_admin():
    if not BOOTSTRAP_ADMIN_EMAIL or not BOOTSTRAP_ADMIN_PASSWORD:
        return
    validate_password(BOOTSTRAP_ADMIN_PASSWORD)
    with SessionLocal() as db:
        if not db.scalar(select(User).where(User.email == BOOTSTRAP_ADMIN_EMAIL)):
            db.add(User(matricule=BOOTSTRAP_ADMIN_MATRICULE, full_name=BOOTSTRAP_ADMIN_NAME,
                        email=BOOTSTRAP_ADMIN_EMAIL, department="Systèmes d'information", site="DIASS",
                        role="admin", password_hash=pwd_context.hash(BOOTSTRAP_ADMIN_PASSWORD),
                        must_change_password=True))
            db.commit()


bootstrap_admin()


@app.exception_handler(401)
async def unauthorized(request: Request, exc):
    return RedirectResponse("/login", status_code=303)


@app.get("/health")
def health(db: Session = Depends(db_session)):
    db.scalar(select(func.count(User.id)))
    return {"status": "ok", "application": APP_NAME}


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, db: Session = Depends(db_session)):
    if optional_user(request, db):
        return RedirectResponse("/", status_code=303)
    return templates.TemplateResponse("login.html", context(request, user=None))


@app.post("/login")
def login(email: str = Form(...), password: str = Form(...), db: Session = Depends(db_session)):
    user = db.scalar(select(User).where(User.email == email.strip().lower()))
    if not user or not user.is_active or not pwd_context.verify(password, user.password_hash):
        return flash_redirect("/login", "Identifiants incorrects", "danger")
    session_data = {"uid": user.id, "csrf": secrets.token_urlsafe(32)}
    response = RedirectResponse("/profil/mot-de-passe" if user.must_change_password else "/", status_code=303)
    response.set_cookie("session", serializer.dumps(session_data), httponly=True, secure=COOKIE_SECURE,
                        samesite="lax", max_age=SESSION_MAX_AGE, path="/")
    return response


@app.post("/logout")
def logout(request: Request, csrf: str = Form(...)):
    verify_csrf(request, csrf)
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie("session", path="/")
    return response


@app.get("/profil/mot-de-passe", response_class=HTMLResponse)
def password_page(request: Request, user: User = Depends(current_user)):
    return templates.TemplateResponse("password.html", context(request, user=user))


@app.post("/profil/mot-de-passe")
def change_password(request: Request, current_password: str = Form(...), new_password: str = Form(...), confirm_password: str = Form(...), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    if not pwd_context.verify(current_password, user.password_hash):
        return flash_redirect("/profil/mot-de-passe", "Mot de passe actuel incorrect", "danger")
    if new_password != confirm_password:
        return flash_redirect("/profil/mot-de-passe", "Les nouveaux mots de passe ne correspondent pas", "danger")
    validate_password(new_password)
    user.password_hash = pwd_context.hash(new_password)
    user.must_change_password = False
    db.commit()
    return flash_redirect("/", "Mot de passe modifié")


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, q: str = "", status: str = "", site: str = "", user: User = Depends(current_user), db: Session = Depends(db_session)):
    stmt = select(Absence).join(User, Absence.agent_id == User.id).order_by(Absence.submitted_at.desc())
    if user.role == "agent":
        stmt = stmt.where(Absence.agent_id == user.id)
    elif user.role == "manager":
        stmt = stmt.where(User.department == user.department, User.site == user.site)
    if status:
        stmt = stmt.where(Absence.status == status)
    if site:
        stmt = stmt.where(User.site == site)
    if q.strip():
        term = f"%{q.strip()}%"
        stmt = stmt.where(or_(User.full_name.ilike(term), User.matricule.ilike(term), Absence.absence_type.ilike(term)))
    absences = db.scalars(stmt.limit(300)).all()
    base_stmt = select(Absence)
    if user.role == "agent":
        base_stmt = base_stmt.where(Absence.agent_id == user.id)
    elif user.role == "manager":
        base_stmt = base_stmt.join(User).where(User.department == user.department, User.site == user.site)
    all_scope = db.scalars(base_stmt).all()
    counts = {s: sum(1 for a in all_scope if a.status == s) for s in ["soumis", "complement_requis", "en_cours", "valide", "rejete"]}
    return templates.TemplateResponse("dashboard.html", context(request, user=user, absences=absences, counts=counts, q=q, selected_status=status, selected_site=site))


@app.get("/absences/nouvelle", response_class=HTMLResponse)
def new_absence_page(request: Request, user: User = Depends(current_user)):
    return templates.TemplateResponse("new_absence.html", context(request, user=user, max_upload_mb=MAX_UPLOAD_MB))


@app.post("/absences/nouvelle")
def create_absence(request: Request, absence_type: str = Form(...), start_date: date = Form(...), end_date: date = Form(...), reason: str = Form(""), documents: list[UploadFile] = File(default=[]), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    if end_date < start_date:
        raise HTTPException(400, "La date de fin ne peut pas précéder la date de début")
    if (end_date - start_date).days > 365:
        raise HTTPException(400, "La période ne peut pas dépasser 366 jours")
    absence = Absence(agent_id=user.id, absence_type=absence_type[:80], start_date=start_date, end_date=end_date, reason=reason.strip()[:4000])
    db.add(absence); db.flush()
    try:
        for doc in documents:
            if doc.filename:
                save_upload(db, absence, user, doc)
        add_event(db, absence, user, "SOUMISSION", "Dossier soumis par l'agent")
        db.commit()
    except Exception:
        db.rollback()
        for path in UPLOAD_DIR.glob(f"{absence.id}_*"):
            path.unlink(missing_ok=True)
        raise
    return flash_redirect(f"/absences/{absence.id}", "Dossier soumis avec succès")


@app.get("/absences/{absence_id}", response_class=HTMLResponse)
def absence_detail(absence_id: int, request: Request, user: User = Depends(current_user), db: Session = Depends(db_session)):
    absence = db.get(Absence, absence_id)
    if not absence or not can_access(user, absence):
        raise HTTPException(404)
    return templates.TemplateResponse("detail.html", context(request, user=user, absence=absence))


@app.post("/absences/{absence_id}/documents")
def add_documents(absence_id: int, request: Request, documents: list[UploadFile] = File(...), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    absence = db.get(Absence, absence_id)
    if not absence or not can_access(user, absence): raise HTTPException(404)
    if user.id == absence.agent_id and absence.status in {"valide", "rejete"}: raise HTTPException(400, "Ce dossier est clôturé")
    count = 0
    for doc in documents:
        if doc.filename: save_upload(db, absence, user, doc); count += 1
    if absence.status == "complement_requis" and absence.agent_id == user.id: absence.status = "en_cours"
    add_event(db, absence, user, "AJOUT_DOCUMENT", f"{count} document(s) ajouté(s)")
    db.commit()
    return flash_redirect(f"/absences/{absence_id}", "Document(s) ajouté(s)")


@app.get("/documents/{document_id}")
def download_document(document_id: int, user: User = Depends(current_user), db: Session = Depends(db_session)):
    doc = db.get(Document, document_id)
    if not doc or not can_access(user, doc.absence): raise HTTPException(404)
    path = UPLOAD_DIR / doc.stored_name
    if not path.exists(): raise HTTPException(404, "Fichier introuvable")
    return FileResponse(path, media_type=doc.content_type, filename=doc.original_name, headers={"X-Content-Type-Options": "nosniff", "Cache-Control": "private, no-store"})


@app.post("/absences/{absence_id}/message")
def add_message(absence_id: int, request: Request, body: str = Form(...), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    absence = db.get(Absence, absence_id)
    if not absence or not can_access(user, absence): raise HTTPException(404)
    body = body.strip()[:4000]
    if not body: raise HTTPException(400, "Message vide")
    db.add(Message(absence_id=absence.id, author_id=user.id, body=body, kind="reponse_agent" if user.id == absence.agent_id else "commentaire_rh"))
    add_event(db, absence, user, "MESSAGE", "Nouveau message dans le dossier")
    db.commit()
    return flash_redirect(f"/absences/{absence_id}", "Message envoyé")


@app.post("/absences/{absence_id}/complement")
def request_complement(absence_id: int, request: Request, body: str = Form(...), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    if user.role not in {"rh", "admin"}: raise HTTPException(403)
    absence = db.get(Absence, absence_id)
    if not absence: raise HTTPException(404)
    body = body.strip()[:4000]
    if not body: raise HTTPException(400, "Précisez le complément attendu")
    absence.status = "complement_requis"
    db.add(Message(absence_id=absence.id, author_id=user.id, body=body, kind="demande_complement"))
    add_event(db, absence, user, "COMPLEMENT_REQUIS", body)
    db.commit()
    send_email(absence.agent.email, f"Complément requis — dossier #{absence.id:05d}", f"Bonjour {absence.agent.full_name},\n\nUn complément est requis :\n{body}\n\nConnectez-vous à JustifAbs.")
    return flash_redirect(f"/absences/{absence_id}", "Demande de complément envoyée")


@app.post("/absences/{absence_id}/decision")
def decision(absence_id: int, request: Request, decision: str = Form(...), note: str = Form(""), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    if user.role not in {"rh", "admin"}: raise HTTPException(403)
    if decision not in {"valide", "rejete", "en_cours"}: raise HTTPException(400)
    absence = db.get(Absence, absence_id)
    if not absence: raise HTTPException(404)
    absence.status = decision; absence.decision_note = note.strip()[:4000]; absence.validated_by = user.id
    absence.validated_at = datetime.utcnow() if decision in {"valide", "rejete"} else None
    add_event(db, absence, user, decision.upper(), absence.decision_note)
    db.commit()
    send_email(absence.agent.email, f"Mise à jour du dossier #{absence.id:05d}", f"Bonjour {absence.agent.full_name},\n\nVotre dossier est désormais : {decision.replace('_',' ')}.\n{absence.decision_note}")
    return flash_redirect(f"/absences/{absence_id}", "Décision enregistrée")


@app.get("/admin/utilisateurs", response_class=HTMLResponse)
def users_page(request: Request, user: User = Depends(current_user), db: Session = Depends(db_session)):
    if user.role != "admin": raise HTTPException(403)
    users = db.scalars(select(User).order_by(User.full_name)).all()
    return templates.TemplateResponse("users.html", context(request, user=user, users=users))


@app.post("/admin/utilisateurs")
def create_user(request: Request, matricule: str = Form(...), full_name: str = Form(...), email: str = Form(...), department: str = Form(""), site: str = Form("DIASS"), role: str = Form("agent"), password: str = Form(...), csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    if user.role != "admin": raise HTTPException(403)
    if role not in {"agent", "manager", "rh", "admin"}: raise HTTPException(400)
    validate_password(password)
    email = email.strip().lower(); matricule = matricule.strip().upper()
    if db.scalar(select(User).where(or_(User.email == email, User.matricule == matricule))):
        return flash_redirect("/admin/utilisateurs", "E-mail ou matricule déjà utilisé", "danger")
    db.add(User(matricule=matricule, full_name=full_name.strip()[:150], email=email, department=department.strip()[:150], site=site.strip()[:100], role=role, password_hash=pwd_context.hash(password), must_change_password=True))
    db.commit()
    return flash_redirect("/admin/utilisateurs", "Utilisateur créé")


@app.post("/admin/utilisateurs/{user_id}/toggle")
def toggle_user(user_id: int, request: Request, csrf: str = Form(...), user: User = Depends(current_user), db: Session = Depends(db_session)):
    verify_csrf(request, csrf)
    if user.role != "admin": raise HTTPException(403)
    target = db.get(User, user_id)
    if not target: raise HTTPException(404)
    if target.id == user.id: return flash_redirect("/admin/utilisateurs", "Vous ne pouvez pas désactiver votre propre compte", "danger")
    target.is_active = not target.is_active; db.commit()
    return flash_redirect("/admin/utilisateurs", "Statut du compte modifié")


def report_rows(db: Session):
    return db.scalars(select(Absence).order_by(Absence.submitted_at.desc())).all()


@app.get("/rapports/export.csv")
def export_csv(user: User = Depends(current_user), db: Session = Depends(db_session)):
    if user.role not in {"rh", "admin"}: raise HTTPException(403)
    output = io.StringIO(); writer = csv.writer(output, delimiter=";")
    writer.writerow(["ID", "Matricule", "Agent", "Département", "Site", "Type", "Début", "Fin", "Jours", "Statut", "Soumis le"])
    for a in report_rows(db): writer.writerow([a.id, a.agent.matricule, a.agent.full_name, a.agent.department, a.agent.site, a.absence_type, a.start_date, a.end_date, (a.end_date-a.start_date).days+1, a.status, a.submitted_at.isoformat(timespec="minutes")])
    data = "\ufeff" + output.getvalue()
    return StreamingResponse(iter([data]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=justificatifs_absences.csv"})


@app.get("/rapports/export.xlsx")
def export_xlsx(user: User = Depends(current_user), db: Session = Depends(db_session)):
    if user.role not in {"rh", "admin"}: raise HTTPException(403)
    wb = Workbook(); ws = wb.active; ws.title = "Absences"
    ws.append(["ID", "Matricule", "Agent", "Département", "Site", "Type", "Début", "Fin", "Jours", "Statut", "Soumis le"])
    for a in report_rows(db): ws.append([a.id, a.agent.matricule, a.agent.full_name, a.agent.department, a.agent.site, a.absence_type, a.start_date, a.end_date, (a.end_date-a.start_date).days+1, a.status, a.submitted_at])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 35)
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=justificatifs_absences.xlsx"})


@app.get("/rapports/export.pdf")
def export_pdf(user: User = Depends(current_user), db: Session = Depends(db_session)):
    if user.role not in {"rh", "admin"}: raise HTTPException(403)
    stream = io.BytesIO(); c = canvas.Canvas(stream, pagesize=landscape(A4)); width, height = landscape(A4)
    def header():
        c.setFont("Helvetica-Bold", 14); c.drawString(35, height-35, "AIBD SA — Rapport des justificatifs d'absence")
        c.setFont("Helvetica", 8); c.drawRightString(width-35, height-35, datetime.now().strftime("Édité le %d/%m/%Y à %H:%M"))
    header(); y = height-65; c.setFont("Helvetica-Bold", 7)
    columns = [(35,"Réf."),(75,"Matricule"),(135,"Agent"),(285,"Site"),(335,"Type"),(465,"Période"),(570,"Statut")]
    for x, label in columns: c.drawString(x,y,label)
    y -= 14; c.setFont("Helvetica", 7)
    for a in report_rows(db):
        if y < 35: c.showPage(); header(); y = height-65
        values = [f"#{a.id:05d}", a.agent.matricule, a.agent.full_name[:25], a.agent.site, a.absence_type[:20], f"{a.start_date:%d/%m/%Y} - {a.end_date:%d/%m/%Y}", a.status.replace('_',' ')]
        for (x,_), value in zip(columns, values): c.drawString(x,y,str(value))
        y -= 13
    c.save(); stream.seek(0)
    return StreamingResponse(stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=justificatifs_absences.pdf"})
